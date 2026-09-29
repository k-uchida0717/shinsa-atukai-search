"""支払基金「審査の一般的な取扱い（医科）」の一覧Excelを取得し data/atukai.json を生成する。

公開ページから Excel のリンクを毎回たどるので、ファイル名が変わっても追従できる。
データ内容に変化がなければ JSON を書き換えない（不要なコミットを防ぐ）。
"""
from __future__ import annotations

import datetime as dt
import io
import json
import os
import re
import sys
import urllib.request
from pathlib import Path
from urllib.parse import urljoin

import openpyxl

PAGE_URL = "https://www.ssk.or.jp/shinryohoshu/sinsa_jirei/kikin_shinsa_atukai/shinsa_atukai_i/index.html"
OUT = Path(os.environ.get("ATUKAI_OUT") or Path(__file__).resolve().parent.parent / "data" / "atukai.json")
UA = {"User-Agent": "Mozilla/5.0 (shinsa-atukai-search; +https://github.com/)"}


def fetch(url: str) -> bytes:
    req = urllib.request.Request(url, headers=UA)
    with urllib.request.urlopen(req, timeout=60) as r:
        return r.read()


def to_date(v) -> str:
    if v is None or v == "":
        return ""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    s = str(v).strip()
    if re.fullmatch(r"\d{5}(\.0)?", s):  # Excel シリアル値
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(float(s)))).isoformat()
    return s


def clean(v) -> str:
    if v is None:
        return ""
    return str(v).replace("\r\n", "\n").strip()


def fix_doubled_title(t: str) -> str:
    """元Excelで「AAB B」のように文言が二重入力されたタイトルを「AB」に戻す。"""
    n = len(t)
    for a in range(4, n // 2 + 1):
        if t[:a] == t[a:2 * a]:
            rest = t[2 * a:]
            if rest and len(rest) % 2 == 0 and rest[:len(rest) // 2] == rest[len(rest) // 2:]:
                return t[:a] + rest[:len(rest) // 2]
    return t


def status_of(*texts: str) -> tuple[str, str]:
    """「R8.6.30削除」「R8.6.1更新」などの注記から状態と日付を取り出す。"""
    found = [(m.group(4), f"令和{m.group(1)}年{m.group(2)}月{m.group(3)}日")
             for t in texts for m in re.finditer(r"Ｒ?R?(\d+)\.(\d+)\.(\d+)\s*(削除|更新)", t)]
    for kind in ("削除", "更新"):
        for k, d in found:
            if k == kind:
                return kind, d
    return "", ""


def find_header(ws, required: str):
    for i, row in enumerate(ws.iter_rows(values_only=True), start=1):
        cells = [clean(c).replace("\n", "") for c in row]
        if required in cells:
            return i, cells
    raise ValueError(f"見出し行（{required}）が見つかりません: {ws.title}")


def col(header: list[str], *names: str) -> int | None:
    for n in names:
        for i, h in enumerate(header):
            if h == n:
                return i
    return None


def parse_ichiran(data: bytes) -> tuple[list[dict], str]:
    ws = openpyxl.load_workbook(io.BytesIO(data), read_only=True).active
    as_of = ""
    for row in ws.iter_rows(max_row=4, values_only=True):
        for c in row:
            m = re.search(r"\d{4}年\s*\d{1,2}月\s*時点", str(c or ""))
            if m and not as_of:
                as_of = m.group(0)
    hr, h = find_header(ws, "タイトル")
    ix = {k: col(h, *names) for k, names in {
        "no": ("№", "No", "NO"), "kubun_no": ("区分Ｎｏ", "区分No"), "kubun": ("区分", "診療項目"),
        "title": ("タイトル",), "rule": ("取扱い",), "basis": ("取扱いを作成した根拠等",),
        "date": ("公表年月日",), "kai": ("公表回",),
    }.items()}
    items = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        g = lambda k: row[ix[k]] if ix[k] is not None and ix[k] < len(row) else None
        raw_title = clean(g("title"))
        if not raw_title:
            continue
        rule, basis = clean(g("rule")), clean(g("basis"))
        status, status_date = status_of(raw_title, rule, basis)
        # 「R8.6.30削除」のような注記行はタイトル表示から外す（status に保持）
        title = re.sub(r"^Ｒ?R?\d+\.\d+\.\d+\s*(削除|更新)\s*", "", raw_title).replace("\n", " ").strip()
        items.append({
            "id": f"s{clean(g('no'))}",
            "src": "支払基金",
            "no": clean(g("no")),
            "kubunNo": clean(g("kubun_no")),
            "kubun": clean(g("kubun")),
            "title": fix_doubled_title(title),
            "rule": rule,
            "basis": basis,
            "date": to_date(g("date")),
            "kai": clean(g("kai")),
            "status": status,
            "statusDate": status_date,
        })
    return items, as_of


def parse_kokuho(data: bytes) -> list[dict]:
    ws = openpyxl.load_workbook(io.BytesIO(data), read_only=True).active
    hr, h = find_header(ws, "タイトル")
    ix = {"no": col(h, "№"), "kubun": col(h, "診療項目", "区分"), "title": col(h, "タイトル"),
          "rule": col(h, "取扱い"), "date": col(h, "公表年月日")}
    items = []
    for row in ws.iter_rows(min_row=hr + 1, values_only=True):
        g = lambda k: row[ix[k]] if ix[k] is not None and ix[k] < len(row) else None
        if not clean(g("title")):
            continue
        items.append({
            "id": f"k{clean(g('no'))}",
            "src": "国保中央会合意",
            "no": clean(g("no")),
            "kubunNo": "",
            "kubun": clean(g("kubun")),
            "title": clean(g("title")),
            "rule": clean(g("rule")),
            "basis": "",
            "date": to_date(g("date")),
            "kai": "",
            "status": "",
            "statusDate": "",
        })
    return items


def verify(html: str, items: list[dict]) -> list[str]:
    """公開ページに記載の総事例数・各回の事例数と Excel の内容を突き合わせる。"""
    problems = []
    total = re.search(r"診療項目順（(\d+)事例）", html)
    if total and int(total.group(1)) != len(items):
        problems.append(f"総事例数: ページ記載 {total.group(1)} / Excel {len(items)}")
    per_kai: dict[str, int] = {}
    for i in items:
        per_kai[i["kai"]] = per_kai.get(i["kai"], 0) + 1
    for kai, n in re.findall(r"（(第\d+回)）」（(\d+)事例）", html):
        if per_kai.get(kai, 0) != int(n):
            problems.append(f"{kai}: ページ記載 {n} / Excel {per_kai.get(kai, 0)}")
    nos = sorted(int(i["no"]) for i in items if i["no"].isdigit())
    if nos != list(range(1, len(items) + 1)):
        problems.append("№ が 1 から連番になっていません")
    for i in items:
        if not i["rule"]:
            problems.append(f"№{i['no']} 取扱いが空です")
    return problems


def main() -> int:
    html = fetch(PAGE_URL).decode("utf-8", errors="replace")
    links = re.findall(r'href="([^"]+\.xlsx)"', html)
    ichiran = next((l for l in links if "ichiran" in l), None)
    kokuho = next((l for l in links if "kokuho" in l), None)
    if not ichiran:
        print("一覧Excelのリンクが見つかりません", file=sys.stderr)
        return 1

    items, as_of = parse_ichiran(fetch(urljoin(PAGE_URL, ichiran)))
    if len(items) < 100:  # 取得失敗・形式変更の安全弁
        print(f"件数が少なすぎます({len(items)})。形式変更の可能性があるため中止", file=sys.stderr)
        return 1
    problems = verify(html, items)
    if problems:
        print("整合性チェックに失敗したため更新を中止:\n  " + "\n  ".join(problems), file=sys.stderr)
        return 1
    if kokuho:
        items += parse_kokuho(fetch(urljoin(PAGE_URL, kokuho)))

    latest = max((i for i in items if i["kai"]), key=lambda i: i["date"], default=None)
    payload = {
        "source": PAGE_URL,
        "asOf": as_of,
        "latestKai": latest["kai"] if latest else "",
        "latestDate": latest["date"] if latest else "",
        "count": len(items),
        "items": items,
    }

    if OUT.exists():
        old = json.loads(OUT.read_text(encoding="utf-8"))
        if old.get("items") == items and old.get("asOf") == as_of:
            print(f"変更なし（{len(items)}件, {as_of}）")
            return 0
    payload["updatedAt"] = dt.datetime.now(dt.timezone(dt.timedelta(hours=9))).isoformat(timespec="minutes")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(payload, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
    print(f"更新しました: {len(items)}件, {as_of}, 最新 {payload['latestKai']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
